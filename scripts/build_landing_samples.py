"""Create static sample downloads using the repository's installed Office libraries.
No app data, providers or network calls. Run backend/venv/bin/python scripts/build_landing_samples.py.
"""
from pathlib import Path
import json
from html import escape
from docx import Document
from docx.shared import Inches as DocInches, Pt as DocPt, RGBColor as DocColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

ROOT = Path(__file__).resolve().parents[1]
data = json.loads((ROOT / 'frontend/src/components/landing/sample.json').read_text())
out = ROOT / 'frontend/public/landing/samples'
out.mkdir(parents=True, exist_ok=True)
notice = 'DocBuilder sample project. Illustrative briefing, not investment advice. Sources verified 6 October 2026.'

def hyperlink(paragraph, text, url):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True))
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr'); color = OxmlElement('w:color'); color.set(qn('w:val'), '215B60'); props.append(color); run.append(props)
    content = OxmlElement('w:t'); content.text = text; run.append(content); link.append(run); paragraph._p.append(link)

doc = Document()
section = doc.sections[0]; section.page_width = DocInches(8.5); section.page_height = DocInches(11)
section.top_margin = section.bottom_margin = DocInches(.7)
normal = doc.styles['Normal']; normal.font.name = 'Calibri'; normal.font.size = DocPt(11); normal.paragraph_format.space_after = DocPt(8)
for name in ['Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.color.rgb = DocColor.from_string('222222')
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
doc.add_paragraph(data['title'], 'Title'); doc.add_paragraph(data['subtitle'])
for item in data['sections']:
    doc.add_heading(item['title'], 1); doc.add_paragraph(item['text']).paragraph_format.keep_with_next = True
    for j, bullet in enumerate(item['bullets']):
        paragraph = doc.add_paragraph(bullet, 'List Bullet'); paragraph.paragraph_format.keep_with_next = j < len(item['bullets']) - 1
doc.add_heading('References', 1)
for i, source in enumerate(data['sources'], 1):
    paragraph = doc.add_paragraph(f'[{i}] {source["publisher"]} '); hyperlink(paragraph, source['title'], source['url'])
footer = section.footer.paragraphs[0]; footer.text = 'DocBuilder AI  /  Sample project'; footer.style = 'Caption'
doc.core_properties.title = data['title']; doc.core_properties.subject = 'Source-grounded landing demonstration'; doc.core_properties.author = 'DocBuilder AI'
doc.save(out / 'clean-energy-outlook.docx')

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
def box(slide, x,y,w,h, text, size, color='223E30', bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=shape.text_frame; tf.word_wrap=True
    for i,line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.font.name='Calibri'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color); p.space_after=Pt(12)
    return shape
for i, item in enumerate(data['slides']):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor.from_string('F6F4EB')
    box(slide,.7,.5,11,.5,item['eyebrow'].upper(),14)
    box(slide,.7,1.4,7,2,item['title'],36,bold=True)
    box(slide,.7,3.7,7,2.5,item['body'],23)
    # Editable, schematic energy landscape, not a chart or photographic evidence.
    shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.7), Inches(1.4), Inches(3.8), Inches(4.8)); shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string('D9E4CF'); shape.line.fill.background()
    for x,y in [(9.5,3.4),(11,2.5)]:
        line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x), Inches(5.5)); line.line.color.rgb=RGBColor.from_string('56705F'); line.line.width=Pt(3)
        for dx,dy in [(0,-.7),(.6,.3),(-.6,.3)]:
            line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x+dx), Inches(y+dy)); line.line.color.rgb=RGBColor.from_string('56705F'); line.line.width=Pt(3)
    box(slide,8.9,5.7,3.4,.4,'Schematic illustration',10)
    if i == len(data['slides']) - 1:
        for j, source in enumerate(data['sources']):
            linkbox=box(slide,.7,5.8+j*.3,7,.35,source['url'],11)
            linkbox.text_frame.paragraphs[0].runs[0].hyperlink.address=source['url']
    box(slide,.7,6.7,11.8,.45,f'DocBuilder sample presentation                                      {i+1} / {len(data["slides"])}',12)
    notes='Source-grounded sample. Separately authored presentation, not instant report conversion.\n'
    notes+='\n'.join(f'{s["title"]}: {s["url"]}' for s in data['sources'] if s['id'] in item['sourceIds'])
    slide.notes_slide.notes_text_frame.text=notes
prs.core_properties.title=data['title']+' leadership presentation'; prs.core_properties.author='DocBuilder AI'
prs.save(out / 'clean-energy-briefing.pptx')

md=f'# {data["title"]}\n\n{data["subtitle"]}\n\n'
html=f'<h1>{escape(data["title"])}</h1><p>{escape(data["subtitle"])}</p>'
for item in data['sections']:
    md+=f'## {item["title"]}\n\n{item["text"]}\n\n'+''.join(f'- {b}\n' for b in item['bullets'])+'\n'
    html+=f'<h2>{escape(item["title"])}</h2><p>{escape(item["text"])}</p><ul>'+''.join(f'<li>{escape(b)}</li>' for b in item['bullets'])+'</ul>'
md+='## References\n\n'; html+='<h2>References</h2><ol>'
for source in data['sources']:
    md+=f'- [{source["title"]}]({source["url"]}) — {source["publisher"]}\n'
    html+=f'<li><a href="{source["url"]}">{escape(source["title"])}</a> — {escape(source["publisher"])}</li>'
html+='</ol><footer>'+notice+'</footer>'
(out / 'clean-energy-outlook.md').write_text(md+'\n'+notice+'\n')
(out / 'clean-energy-outlook.html').write_text('<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Clean Energy Outlook — Sample</title><style>body{max-width:720px;margin:48px auto;padding:0 24px;font:17px/1.6 Georgia,serif;color:#223128;background:#fff}h1{font-size:40px}h2{font-size:24px;margin-top:30px}a{color:#215b60}footer{font:12px/1.5 sans-serif;border-top:1px solid #ddd;padding-top:16px}@media print{body{margin:0}a{color:inherit}}</style><body>'+html+'<script>if(new URLSearchParams(location.search).has("print"))addEventListener("load",()=>print())</script></body></html>')
print('Built DOCX, PPTX, Markdown and HTML samples.')
